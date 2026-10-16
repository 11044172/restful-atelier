"""Literal XLSX output and bounded import, preserving strings without CSV escapes."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from django.core.management.base import CommandError
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils.exceptions import IllegalCharacterError

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 100


class ProductDataWriter:
    def __init__(self):
        self.workbook = Workbook(write_only=True)
        self.sheet = self.workbook.create_sheet("Products")

    def writerow(self, values):
        cells = []
        for value in values:
            text = "" if value is None else str(value)
            if len(text) > 32767:
                raise ValueError("商品欄位超過 XLSX 的 32,767 字元上限，請先縮短內容。")
            try:
                cell = WriteOnlyCell(self.sheet, value=text)
            except IllegalCharacterError as exc:
                raise ValueError("商品欄位包含無法匯出的控制字元。") from exc
            # openpyxl initially infers formulas for '=...'; explicitly override it.
            cell.data_type = "s"
            cells.append(cell)
        self.sheet.append(cells)

    def content(self):
        output = BytesIO()
        self.workbook.save(output)
        # XML parsers normalize raw CR characters to LF. Preserve literal text
        # even with the stdlib XML writer used when lxml is not installed.
        preserved = BytesIO()
        with ZipFile(output) as source, ZipFile(preserved, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data = source.read(entry)
                if entry.filename.startswith("xl/worksheets/"):
                    data = data.replace(b"\r", b"&#13;")
                target.writestr(entry, data)
        return preserved.getvalue()


def read_product_workbook(path):
    try:
        if path.stat().st_size > MAX_EXPANDED_BYTES:
            raise ValueError("XLSX file is too large")
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
            ):
                raise ValueError("XLSX expanded data is too large")
        workbook = load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("Use a workbook with exactly one worksheet")
            sheet = workbook.worksheets[0]
            # Ignore untrusted dimension hints; enforce limits while streaming cells.
            sheet.reset_dimensions()
            rows = []
            headers = None
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                if row_number > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise ValueError("XLSX has too many rows or columns")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Formulas are not accepted; format cells as text")
                values = [
                    "" if cell.value is None else str(cell.value) for cell in cells
                ]
                if headers is None:
                    headers = values
                    if len(headers) != len(set(headers)):
                        raise ValueError("Duplicate column headers")
                elif any(values):
                    if len(values) > len(headers) and any(values[len(headers) :]):
                        raise ValueError("Data outside the header columns")
                    rows.append(dict(zip(headers, values)))
            return headers or [], rows
        finally:
            workbook.close()
    except CommandError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, TypeError) as exc:
        raise CommandError(f"Invalid XLSX: {exc}") from exc

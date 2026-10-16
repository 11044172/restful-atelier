"""Regression checks for SEC-01 through SEC-09; no external writes or messages."""

import logging
import sys
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Barrier
from unittest.mock import Mock, patch
from uuid import uuid4

import boto3
from botocore.config import Config
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core import mail
from django.core.cache import cache
from django.core.exceptions import PermissionDenied
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connections
from django.test import (
    Client,
    RequestFactory,
    TestCase,
    TransactionTestCase,
    override_settings,
    skipUnlessDBFeature,
)
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from PIL import Image

from catalog import product_image_service as product_images
from catalog.models import Product, ProductCategory, ProductImage
from catalog.product_data_files import ProductDataWriter, read_product_workbook
from content import project_image_service as project_images
from core import direct_image_uploads as single_images
from core.admin_image_uploads import generate_presigned_put
from core.admin_site import backoffice_site
from core.logging_filters import SensitivePathFormatter, redact_signed_paths
from core.middleware import RequestObservabilityMiddleware, safe_request_path
from core.models import SiteSettings
from orders.forms import CheckoutForm
from orders.line_login import LineLoginError
from orders.models import (
    LineCustomer,
    LineNotification,
    NotificationOutbox,
    Order,
    OrderAuditLog,
    Payment,
    PaymentMethod,
)
from orders.notifications import send_order_email_event


@override_settings(
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    ECPAY_INVOICE_ENABLED=False,
)
class SecurityRemediationTests(TestCase):
    def staff(self, *permissions):
        user = get_user_model().objects.create_user(
            f"staff-{uuid4().hex[:10]}", is_staff=True
        )
        user.user_permissions.add(
            *Permission.objects.filter(
                content_type__app_label="orders", codename__in=permissions
            )
        )
        return user

    def order(self):
        return Order.objects.create(
            idempotency_key=uuid4().hex,
            customer_name="Synthetic customer",
            email="customer@example.invalid",
            phone="000",
            shipping_information="Synthetic address",
            subtotal=Decimal("1000"),
            shipping_fee=0,
            status=Order.Status.AWAITING_PAYMENT,
        )

    def action(self, user, model, action, pk):
        client = Client(enforce_csrf_checks=True)
        client.force_login(user)
        url = reverse(f"admin:orders_{model}_changelist")
        page = client.get(url)
        self.assertEqual(page.status_code, 200)
        response = client.post(
            url,
            {
                "action": action,
                "_selected_action": str(pk),
                "index": "0",
                "csrfmiddlewaretoken": client.cookies["csrftoken"].value,
            },
        )
        self.assertIn(response.status_code, (200, 302))
        return page

    def payment(self):
        method, _ = PaymentMethod.objects.get_or_create(
            code="bank_transfer", defaults={"display_name": "Audit"}
        )
        return Payment.objects.create(
            order=self.order(),
            method=method,
            amount=1000,
            status=Payment.Status.CONFIRMED,
            refund_reason="Authorized refund reason",
        )

    def test_viewer_cannot_refund_even_with_valid_csrf_and_direct_post(self):
        payment = self.payment()
        page = self.action(
            self.staff("view_payment"),
            "payment",
            "record_remaining_full_refund",
            payment.pk,
        )
        self.assertNotContains(page, 'value="record_remaining_full_refund"')
        payment.refresh_from_db()
        payment.order.refresh_from_db()
        self.assertEqual(payment.refunded_amount, 0)
        self.assertEqual(payment.status, Payment.Status.CONFIRMED)
        self.assertEqual(payment.order.status, Order.Status.PAID)
        self.assertFalse(
            OrderAuditLog.objects.filter(
                order=payment.order, event="refund_completed"
            ).exists()
        )

    def test_refund_requires_all_three_permissions_and_records_authorized_actor(self):
        required = {"change_payment", "change_order", "record_refund"}
        for missing in required:
            payment = self.payment()
            self.action(
                self.staff("view_payment", *(required - {missing})),
                "payment",
                "record_remaining_full_refund",
                payment.pk,
            )
            payment.refresh_from_db()
            self.assertEqual(payment.refunded_amount, 0)
        payment = self.payment()
        operator = self.staff(*required)
        self.action(operator, "payment", "record_remaining_full_refund", payment.pk)
        payment.refresh_from_db()
        self.assertEqual(payment.refunded_amount, 1000)
        self.assertEqual(payment.refund_operator, operator)
        self.assertTrue(
            OrderAuditLog.objects.filter(
                order=payment.order, actor=operator, event="refund_completed"
            ).exists()
        )

    def test_viewer_cannot_requeue_outbox_or_line_notification(self):
        order = self.order()
        job = NotificationOutbox.objects.create(
            order=order,
            channel="email",
            event_type="order_received",
            dedupe_key=uuid4().hex,
            status="dead",
        )
        page = self.action(
            self.staff("view_notificationoutbox"),
            "notificationoutbox",
            "retry_jobs",
            job.pk,
        )
        self.assertNotContains(page, 'value="retry_jobs"')
        job.refresh_from_db()
        self.assertEqual(job.status, "dead")
        customer = LineCustomer.objects.create(line_user_id="U-notify")
        notification = LineNotification.objects.create(
            order=order,
            line_customer=customer,
            notification_type="order_received",
            dedupe_key=uuid4().hex,
            status="failed",
        )
        page = self.action(
            self.staff("view_linenotification"),
            "linenotification",
            "retry_failed",
            notification.pk,
        )
        self.assertNotContains(page, 'value="retry_failed"')
        self.assertFalse(NotificationOutbox.objects.filter(channel="line").exists())

    def test_authorized_notification_retries_record_operator(self):
        order = self.order()
        job = NotificationOutbox.objects.create(
            order=order,
            channel="email",
            event_type="order_received",
            dedupe_key=uuid4().hex,
            status="dead",
        )
        operator = self.staff("change_notificationoutbox")
        self.action(operator, "notificationoutbox", "retry_jobs", job.pk)
        job.refresh_from_db()
        self.assertEqual(job.status, "pending")
        self.assertTrue(
            OrderAuditLog.objects.filter(
                order=order, actor=operator, event="notification_requeued"
            ).exists()
        )
        notification = LineNotification.objects.create(
            order=order,
            line_customer=LineCustomer.objects.create(line_user_id="U-authorized"),
            notification_type="order_received",
            dedupe_key=uuid4().hex,
            status="failed",
        )
        operator = self.staff("change_linenotification")
        self.action(operator, "linenotification", "retry_failed", notification.pk)
        self.assertTrue(NotificationOutbox.objects.filter(channel="line").exists())
        self.assertTrue(
            OrderAuditLog.objects.filter(
                order=order, actor=operator, event="notification_requeued"
            ).exists()
        )

    def test_action_bodies_enforce_permissions_and_each_object(self):
        payment = self.payment()
        request = RequestFactory().post("/admin/")
        request.user = self.staff("view_payment")
        payment_admin = backoffice_site._registry[Payment]
        with self.assertRaises(PermissionDenied):
            payment_admin.record_remaining_full_refund(
                request, Payment.objects.filter(pk=payment.pk)
            )
        request.user = self.staff("change_payment", "change_order", "record_refund")
        with patch.object(
            payment_admin,
            "has_change_permission",
            side_effect=lambda req, obj=None: obj is None,
        ):
            with self.assertRaises(PermissionDenied):
                payment_admin.record_remaining_full_refund(
                    request, Payment.objects.filter(pk=payment.pk)
                )
        payment.refresh_from_db()
        self.assertEqual(payment.refunded_amount, 0)
        for model, action in (
            (NotificationOutbox, "retry_jobs"),
            (LineNotification, "retry_failed"),
        ):
            request.user = self.staff()
            with self.assertRaises(PermissionDenied):
                getattr(backoffice_site._registry[model], action)(
                    request, model.objects.none()
                )

    @override_settings(
        LINE_LOGIN_CHANNEL_ID="audit",
        LINE_LOGIN_CHANNEL_SECRET="synthetic",
        LINE_LOGIN_CALLBACK_URL="http://testserver/auth/line/callback/",
        LINE_MESSAGING_CHANNEL_ACCESS_TOKEN="synthetic",
        LINE_MESSAGING_CHANNEL_SECRET="synthetic",
        LINE_OFFICIAL_ACCOUNT_BASIC_ID="@audit",
    )
    def test_query_cannot_unblock_customer_or_allow_order_during_line_outage(self):
        SiteSettings.objects.create(
            checkout_enabled=True, line_official_url="https://line.me/R/test"
        )
        product = Product.objects.create(
            name="Audit product",
            slug="audit-order",
            sku="AUDIT-ORDER",
            description="Synthetic",
            price=1000,
            stock=1,
            is_published=True,
        )
        ProductImage.objects.create(product=product, image="products/audit.jpg")
        self.client.post(
            reverse("orders:cart_add", args=[product.slug]), {"quantity": 1}
        )
        customer = LineCustomer.objects.create(
            line_user_id="U-audit", is_friend=False, is_blocked=True
        )
        self.client.get(reverse("line_login_start"))
        oauth = self.client.session["line_oauth"]
        with (
            patch(
                "orders.views.exchange_code",
                return_value={"access_token": "mock", "id_token": "mock"},
            ),
            patch(
                "orders.views.verify_id_token",
                return_value={"sub": "U-audit", "name": "Audit"},
            ),
            patch(
                "orders.views.get_friendship_status",
                side_effect=LineLoginError("outage"),
            ),
        ):
            self.client.get(
                reverse("line_login_callback"),
                {
                    "state": oauth["state"],
                    "code": "mock",
                    "friendship_status_changed": "true",
                },
            )
        customer.refresh_from_db()
        self.assertFalse(customer.is_friend)
        self.assertTrue(customer.is_blocked)
        self.assertFalse(self.client.session["line_friend_verified"])
        self.assertTrue(self.client.session["line_friend_check_failed"])
        self.client.get(reverse("orders:checkout"))
        data = {
            "idempotency_key": self.client.session["checkout_token"],
            "customer_name": "Audit",
            "phone": "0900",
            "email": "audit@example.invalid",
            "shipping_information": "Synthetic address",
            "website": "",
            "policies_accepted": "on",
        }
        form = CheckoutForm(data)
        self.assertTrue(form.is_valid(), form.errors)
        with patch("orders.views.verify_turnstile", return_value=True):
            self.client.post(reverse("orders:checkout"), data)
        self.assertFalse(Order.objects.exists())

    def test_all_signed_paths_are_redacted_in_normal_error_slow_and_exception_logs(
        self,
    ):
        for prefix in ("/pay/", "/shop/cancel/", "/payments/ecpay/return/"):
            path = prefix + "SYNTHETIC-SECRET/"
            self.assertNotIn("SYNTHETIC-SECRET", safe_request_path(path))
            self.assertNotIn(
                "SYNTHETIC-SECRET", redact_signed_paths("Not Found: " + path)
            )
            for code in (200, 404, 500):
                request = RequestFactory().get(path)
                response = Mock(status_code=code)
                with (
                    override_settings(SLOW_REQUEST_MS=0),
                    self.assertLogs("restfull.requests", level="INFO") as captured,
                ):
                    RequestObservabilityMiddleware(lambda req: response)(request)
                self.assertNotIn("SYNTHETIC-SECRET", "\n".join(captured.output))
            try:
                raise ValueError(path)
            except ValueError:
                record = logging.LogRecord(
                    "audit",
                    logging.ERROR,
                    __file__,
                    1,
                    "error %s",
                    (path,),
                    sys.exc_info(),
                )
            self.assertNotIn(
                "SYNTHETIC-SECRET", SensitivePathFormatter().format(record)
            )

    def test_readiness_exposes_only_status_and_caches_migration_graph(self):
        cache.clear()
        with patch("core.views.MigrationExecutor") as executor:
            executor.return_value.migration_plan.return_value = []
            for _ in range(2):
                response = self.client.get("/readiness/")
                self.assertEqual(response.json(), {"status": "ready"})
                self.assertEqual(response["Cache-Control"], "no-store")
            self.assertEqual(executor.call_count, 1)
        cache.clear()
        with patch(
            "core.views.MigrationExecutor", side_effect=RuntimeError("internal-secret")
        ):
            response = self.client.get("/readiness/")
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.json(), {"status": "unavailable"})
        cache.clear()

    @override_settings(
        ORDER_NOTIFICATION_EMAIL="staff@example.invalid",
        ORDER_NOTIFICATION_EMAILS=["internal-ops@example.invalid"],
    )
    def test_customer_email_contains_neither_internal_addresses_nor_admin_url(self):
        order = self.order()
        send_order_email_event(order, "order_received")
        self.assertEqual(len(mail.outbox), 2)
        customer, internal = mail.outbox
        self.assertEqual(customer.to, [order.email])
        self.assertFalse(customer.cc)
        self.assertFalse(customer.bcc)
        for value in (
            "internal-ops@example.invalid",
            "staff@example.invalid",
            "/admin/",
        ):
            self.assertNotIn(value, customer.body)
        self.assertIn("internal-ops@example.invalid", internal.to)
        self.assertIn("/admin/", internal.body)
        self.assertNotIn(order.email, internal.to)

    def test_checkout_and_completion_are_never_cached(self):
        order = self.order()
        session = self.client.session
        session["completed_order"] = order.public_number
        session.save()
        for url in (
            reverse("orders:checkout"),
            reverse("orders:complete", args=[order.public_number]),
        ):
            self.assertIn("no-store", self.client.get(url)["Cache-Control"])

    def test_export_formula_strings_survive_save_and_reimport_without_mutation(self):
        user = get_user_model().objects.create_superuser(
            "export-audit", "audit@example.invalid", "synthetic"
        )
        self.client.force_login(user)
        product = Product.objects.create(
            name="=1+1",
            sku="SAFE-SKU",
            description="@literal\nsecond line",
            category=ProductCategory.objects.create(name="+category", slug="audit"),
            price=100,
        )
        response = self.client.post(
            reverse("admin:catalog_product_changelist"),
            {
                "action": "export_product_data",
                "_selected_action": str(product.pk),
                "index": "0",
            },
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertTrue(
            all(cell.data_type != "f" for row in workbook.active for cell in row)
        )
        self.assertEqual(workbook.active["B2"].value, "=1+1")
        with TemporaryDirectory() as directory:
            path = Path(directory) / "products.xlsx"
            workbook.save(path)
            headers, rows = read_product_workbook(path)
            self.assertEqual(rows[0]["name"], "=1+1")
            call_command(
                "import_product_data", str(path), apply=True, stdout=StringIO()
            )
        product.refresh_from_db()
        self.assertEqual(product.description, "@literal\nsecond line")
        self.assertEqual(product.sku, "SAFE-SKU")

    def test_xlsx_import_rejects_formulas_and_keeps_all_formula_prefixes_literal(self):
        values = [
            "=1+1",
            "+1+1",
            "-1+1",
            "@SUM(1)",
            "\t=1+1",
            "\r=1+1",
            "＝1+1",
            "＋1",
            "－1",
            "＠SUM(1)",
        ]
        writer = ProductDataWriter()
        writer.writerow(values)
        workbook = load_workbook(BytesIO(writer.content()))
        self.assertEqual([cell.value for cell in workbook.active[1]], values)
        self.assertTrue(all(cell.data_type == "s" for cell in workbook.active[1]))
        with TemporaryDirectory() as directory:
            path = Path(directory) / "formula.xlsx"
            workbook = Workbook()
            workbook.active.append(["product_id", "current_sku", "current_description"])
            workbook.active.append([1, "SKU", "=1+1"])
            workbook.save(path)
            with self.assertRaises(CommandError):
                read_product_workbook(path)


@override_settings(
    ADMIN_IMAGE_STAGING_BUCKET="private-staging", ADMIN_IMAGE_UPLOADS_PER_HOUR=60
)
class ImageSecurityRemediationTests(TestCase):
    routes = ("single", "product", "project")
    errors = (
        single_images.DirectImageError,
        product_images.ProductImageError,
        project_images.ProjectImageError,
    )

    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            "image-security", "audit@example.invalid", "synthetic"
        )
        self.storage = Mock()
        self.storage.generate_presigned_url.return_value = (
            "https://storage.example.invalid/put"
        )
        self.objects = {}
        self.storage.head_object.side_effect = lambda **kw: {
            "ContentLength": len(self.objects[(kw["Bucket"], kw["Key"])]),
            "ContentType": "image/jpeg",
        }
        self.storage.get_object.side_effect = lambda **kw: {
            "Body": BytesIO(self.objects[(kw["Bucket"], kw["Key"])])
        }
        self.storage.put_object.side_effect = lambda **kw: self.objects.__setitem__(
            (kw["Bucket"], kw["Key"]), kw["Body"]
        )
        self.storage.delete_object.side_effect = lambda **kw: self.objects.pop(
            (kw["Bucket"], kw["Key"]), None
        )
        stack = self.enterContext(ExitStack())
        for module in (single_images, product_images, project_images):
            stack.enter_context(
                patch.object(
                    module, "_storage_client", return_value=(self.storage, "public")
                )
            )
        stack.enter_context(
            patch(
                "core.admin_image_uploads.storage_connection",
                return_value=(self.storage, "public"),
            )
        )

    @staticmethod
    def raw_image(format="JPEG", *, width=2, height=2, exif=False, animated=False):
        output = BytesIO()
        source = Image.new("RGB", (width, height), "red")
        options = {}
        if exif:
            metadata = Image.Exif()
            metadata[315] = "Sensitive metadata"
            options["exif"] = metadata
        if animated:
            options.update(
                save_all=True,
                append_images=[Image.new("RGB", (width, height), "blue")],
                duration=100,
            )
        source.save(output, format=format, **options)
        return output.getvalue()

    def prepare(self, route, raw, *, width=2, height=2, content_type="image/jpeg"):
        session = uuid4()
        filename = {"image/png": "image.png", "image/webp": "image.webp"}.get(
            content_type, "image.jpg"
        )
        values = dict(
            user=self.user,
            filename=filename,
            content_type=content_type,
            size=len(raw),
            width=width,
            height=height,
        )
        if route == "single":
            image, _ = single_images.create_upload(
                category="site.home_hero_image", **values
            )
            key = image.object_key
            complete = lambda: single_images.complete_upload(
                user=self.user,
                category=image.category,
                upload_id=image.pk,
                object_key=key,
            )[0]
        else:
            service = product_images if route == "product" else project_images
            image, _ = service.create_presigned_upload(upload_session=session, **values)
            key = image.image.name
            complete = lambda: service.complete_upload(
                user=self.user, upload_session=session, object_key=key
            )
        self.objects[("private-staging", key)] = raw
        return image, key, complete

    def test_non_images_are_rejected_on_all_three_paths(self):
        for route in self.routes:
            with self.subTest(route=route):
                image, key, complete = self.prepare(route, b"THIS IS NOT AN IMAGE")
                with self.assertRaises(self.errors):
                    complete()
                self.storage.put_object.assert_not_called()
                self.assertFalse(type(image).objects.filter(pk=image.pk).exists())
                self.assertNotIn(("private-staging", key), self.objects)
        self.assertEqual(self.storage.get_object.call_count, 3)

    def test_actual_dimensions_must_match_declaration_on_all_paths(self):
        for route in self.routes:
            with self.subTest(route=route):
                image, key, complete = self.prepare(route, self.raw_image(width=3))
                with self.assertRaises(self.errors) as raised:
                    complete()
                self.assertEqual(raised.exception.code, "dimensions_mismatch")
        self.storage.put_object.assert_not_called()

    def test_large_dimensions_and_pixel_limits_are_checked_before_decoding(self):
        raw = self.raw_image(width=3, height=3)
        for route in self.routes:
            image, key, complete = self.prepare(route, raw, width=3, height=3)
            with (
                override_settings(ADMIN_IMAGE_MAX_OUTPUT_PIXELS=8),
                self.assertRaises(self.errors) as raised,
            ):
                complete()
            self.assertEqual(raised.exception.code, "dimensions_too_large")
        self.storage.put_object.assert_not_called()

    def test_animated_images_are_rejected_on_all_paths(self):
        raw = self.raw_image("PNG", animated=True)
        self.storage.head_object.side_effect = lambda **kw: {
            "ContentLength": len(raw),
            "ContentType": "image/png",
        }
        for route in self.routes:
            image, key, complete = self.prepare(route, raw, content_type="image/png")
            with self.assertRaises(self.errors) as raised:
                complete()
            self.assertEqual(raised.exception.code, "animated_image")
        self.storage.put_object.assert_not_called()

    def test_valid_images_are_reencoded_without_metadata_and_cannot_be_replaced_by_reput(
        self,
    ):
        raw = self.raw_image(exif=True)
        for route in self.routes:
            with self.subTest(route=route):
                image, staging_key, complete = self.prepare(route, raw)
                verified = complete()
                public_key = (
                    verified.object_key if route == "single" else verified.image.name
                )
                self.assertFalse(public_key.startswith("pending/"))
                self.assertNotEqual(public_key, staging_key)
                clean = self.objects[("public", public_key)]
                self.assertNotEqual(clean, raw)
                with Image.open(BytesIO(clean)) as decoded:
                    self.assertEqual(decoded.size, (2, 2))
                    self.assertFalse(decoded.getexif())
                # A still-valid PUT can recreate staging but cannot address the public snapshot.
                self.storage.put_object(
                    Bucket="private-staging", Key=staging_key, Body=b"replacement"
                )
                self.assertEqual(self.objects[("public", public_key)], clean)
                with self.assertRaises(self.errors):
                    complete()
                self.assertEqual(verified.file_size, len(clean))

    def test_byte_reads_are_bounded_and_stream_is_closed_on_failure(self):
        raw = self.raw_image()
        image, key, complete = self.prepare("single", raw)
        stream = Mock()
        stream.read.return_value = raw + b"extra"
        self.storage.get_object.side_effect = None
        self.storage.get_object.return_value = {"Body": stream}
        with self.assertRaises(single_images.DirectImageError):
            complete()
        from django.conf import settings

        stream.read.assert_called_once_with(settings.ADMIN_IMAGE_MAX_OUTPUT_BYTES + 1)
        stream.close.assert_called_once()
        self.storage.put_object.assert_not_called()

    def test_png_and_webp_alpha_survive_reencoding_on_all_paths(self):
        for format, content_type in (("PNG", "image/png"), ("WEBP", "image/webp")):
            output = BytesIO()
            Image.new("RGBA", (2, 2), (255, 0, 0, 0)).save(output, format=format)
            raw = output.getvalue()
            self.storage.head_object.side_effect = None
            self.storage.head_object.return_value = {
                "ContentLength": len(raw),
                "ContentType": content_type,
            }
            for route in self.routes:
                with self.subTest(format=format, route=route):
                    image, key, complete = self.prepare(
                        route, raw, content_type=content_type
                    )
                    verified = complete()
                    public_key = (
                        verified.object_key
                        if route == "single"
                        else verified.image.name
                    )
                    with Image.open(
                        BytesIO(self.objects[("public", public_key)])
                    ) as decoded:
                        self.assertEqual(decoded.getpixel((0, 0))[3], 0)

    def test_get_snapshot_must_be_valid_even_if_head_was_valid(self):
        raw = self.raw_image()
        image, key, complete = self.prepare("single", raw)
        self.storage.get_object.side_effect = lambda **kw: {
            "Body": BytesIO(b"X" * len(raw))
        }
        with self.assertRaises(single_images.DirectImageError):
            complete()
        self.storage.put_object.assert_not_called()

    def test_expired_pending_uploads_are_rejected_before_storage_access(self):
        raw = self.raw_image()
        for route in self.routes:
            image, key, complete = self.prepare(route, raw)
            type(image).objects.filter(pk=image.pk).update(
                created_at=timezone.now() - timedelta(minutes=11)
            )
            with self.assertRaises(self.errors) as raised:
                complete()
            self.assertEqual(raised.exception.code, "upload_expired")
        self.storage.get_object.assert_not_called()
        self.storage.head_object.assert_not_called()

    def test_shared_issuance_quota_cannot_be_bypassed_by_switching_routes_or_deleting_rows(
        self,
    ):
        raw = self.raw_image()
        with override_settings(ADMIN_IMAGE_UPLOADS_PER_HOUR=1):
            image, key, _ = self.prepare("single", raw)
            image.delete()
            for route in ("product", "project", "single"):
                with self.assertRaises(self.errors) as raised:
                    self.prepare(route, raw)
                self.assertEqual(raised.exception.code, "upload_quota_exceeded")
        self.assertEqual(self.storage.generate_presigned_url.call_count, 1)

    def test_shared_byte_quota_is_enforced(self):
        raw = self.raw_image()
        with override_settings(ADMIN_IMAGE_BYTES_PER_HOUR=len(raw)):
            self.prepare("product", raw)
            with self.assertRaises(project_images.ProjectImageError) as raised:
                self.prepare("project", raw)
            self.assertEqual(raised.exception.status, 429)

    def test_private_bucket_is_required_and_cannot_equal_public_bucket(self):
        raw = self.raw_image()
        for value in ("", "public"):
            with override_settings(ADMIN_IMAGE_STAGING_BUCKET=value):
                with self.assertRaises(single_images.DirectImageError):
                    self.prepare("single", raw)
        self.storage.generate_presigned_url.assert_not_called()

    def test_real_s3_signature_binds_exact_length_and_private_bucket(self):
        from urllib.parse import parse_qs, urlsplit

        client = boto3.client(
            "s3",
            endpoint_url="https://synthetic.r2.cloudflarestorage.com",
            aws_access_key_id="synthetic",
            aws_secret_access_key="synthetic",
            region_name="auto",
            config=Config(signature_version="s3v4"),
        )
        url = generate_presigned_put(
            client=client,
            bucket="public",
            object_key="pending/image.jpg",
            content_type="image/jpeg",
            size=1234,
            expires_in=600,
        )
        parsed = urlsplit(url)
        self.assertIn("private-staging", parsed.netloc + parsed.path)
        self.assertIn(
            "content-length",
            parse_qs(parsed.query)["X-Amz-SignedHeaders"][0].split(";"),
        )


@skipUnlessDBFeature("has_select_for_update")
@override_settings(
    ADMIN_IMAGE_STAGING_BUCKET="private-staging", ADMIN_IMAGE_UPLOADS_PER_HOUR=1
)
class ImageConcurrencySecurityTests(TransactionTestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            "concurrent-images", "audit@example.invalid", "synthetic"
        )
        output = BytesIO()
        Image.new("RGB", (2, 2)).save(output, format="JPEG")
        self.raw = output.getvalue()
        self.storage = Mock()
        self.storage.generate_presigned_url.return_value = (
            "https://storage.example.invalid/put"
        )

    def test_parallel_issuance_on_different_routes_respects_one_shared_quota(self):
        barrier = Barrier(3)

        def attempt(route):
            try:
                barrier.wait(timeout=10)
                values = dict(
                    user=self.user,
                    filename="image.jpg",
                    content_type="image/jpeg",
                    size=len(self.raw),
                    width=2,
                    height=2,
                )
                if route == "single":
                    single_images.create_upload(
                        category="site.home_hero_image", **values
                    )
                else:
                    service = product_images if route == "product" else project_images
                    service.create_presigned_upload(upload_session=uuid4(), **values)
                return "issued"
            except (
                single_images.DirectImageError,
                product_images.ProductImageError,
                project_images.ProjectImageError,
            ) as exc:
                return exc.code
            finally:
                connections.close_all()

        with ExitStack() as stack:
            for module in (single_images, product_images, project_images):
                stack.enter_context(
                    patch.object(
                        module, "_storage_client", return_value=(self.storage, "public")
                    )
                )
            with ThreadPoolExecutor(max_workers=3) as executor:
                results = list(executor.map(attempt, ("single", "product", "project")))
        self.assertEqual(results.count("issued"), 1)
        self.assertEqual(results.count("upload_quota_exceeded"), 2)
        self.assertEqual(self.storage.generate_presigned_url.call_count, 1)

    def test_parallel_completions_publish_exactly_one_snapshot(self):
        self.storage.head_object.return_value = {
            "ContentLength": len(self.raw),
            "ContentType": "image/jpeg",
        }
        self.storage.get_object.side_effect = lambda **kw: {"Body": BytesIO(self.raw)}
        barrier = Barrier(2)
        with patch.object(
            single_images, "_storage_client", return_value=(self.storage, "public")
        ):
            image, _ = single_images.create_upload(
                user=self.user,
                category="site.home_hero_image",
                filename="image.jpg",
                content_type="image/jpeg",
                size=len(self.raw),
                width=2,
                height=2,
            )
            staging_key = image.object_key

            def complete(_):
                try:
                    barrier.wait(timeout=10)
                    single_images.complete_upload(
                        user=self.user,
                        category=image.category,
                        upload_id=image.pk,
                        object_key=staging_key,
                    )
                    return "ready"
                except single_images.DirectImageError as exc:
                    return exc.code
                finally:
                    connections.close_all()

            with ThreadPoolExecutor(max_workers=2) as executor:
                results = list(executor.map(complete, range(2)))
        self.assertCountEqual(results, ["ready", "upload_not_owned"])
        self.storage.put_object.assert_called_once()
        image.refresh_from_db()
        self.assertEqual(image.status, "ready")
        self.assertNotEqual(image.object_key, staging_key)
